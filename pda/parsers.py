# -*- coding: utf-8 -*-
"""按扩展名提取纯文本：txt/md/pdf/docx/xlsx/xlsm/xls/pptx/图片(OCR)。

不限制文件类型：未知扩展名按纯文本尝试（UTF-8/GBK 嗅探）；提取不出文字的
二进制文件抛 BinaryFileError，由 ingest 归档原件并按文件名建索引。

图片 OCR 依赖 rapidocr-onnxruntime，为可选组件：导入失败时图片类型给出
"OCR 组件不可用"提示，不影响其他格式。
"""
import os

SUPPORTED_EXTS = {
    ".txt", ".md", ".pdf", ".docx",
    ".xlsx", ".xlsm", ".xls",
    ".pptx",
    ".png", ".jpg", ".jpeg", ".bmp", ".webp",
}

IMAGE_EXTS = {".png", ".jpg", ".jpeg", ".bmp", ".webp"}


class ParseError(Exception):
    """文件无法解析（无文本内容）。"""


class BinaryFileError(ParseError):
    """二进制文件，提取不出文本（ingest 据此归档原件、按文件名建索引）。"""


def ocr_available() -> bool:
    try:
        import rapidocr_onnxruntime  # noqa: F401

        return True
    except Exception:
        return False


def extract_text(path: str) -> str:
    ext = path.lower().rsplit(".", 1)[-1]
    ext = f".{ext}"
    if ext in (".txt", ".md"):
        return _read_text_file(path)
    if ext == ".pdf":
        return _read_pdf(path)
    if ext == ".docx":
        return _read_docx(path)
    if ext in (".xlsx", ".xlsm"):
        return _read_xlsx(path)
    if ext == ".xls":
        return _read_xls(path)
    if ext == ".pptx":
        return _read_pptx(path)
    if ext in IMAGE_EXTS:
        return _read_image_ocr(path)
    return _read_unknown(path)


# 未知扩展名超过此大小不当文本读（几 GB 的 iso/视频整读进内存会卡死或 MemoryError）
_UNKNOWN_TEXT_MAX = 20 * 1024 * 1024

_UTF16_BOMS = (b"\xff\xfe", b"\xfe\xff")

_BOMS = (
    (b"\xef\xbb\xbf", "utf-8-sig"),
    (b"\xff\xfe", "utf-16"),
    (b"\xfe\xff", "utf-16"),
)


def _decode(raw: bytes) -> str | None:
    """按 BOM → UTF-8 → GBK 顺序解码；都不行返回 None。"""
    for bom, encoding in _BOMS:
        if raw.startswith(bom):
            try:
                return raw.decode(encoding)
            except UnicodeDecodeError:
                return None
    for encoding in ("utf-8", "gbk"):
        try:
            return raw.decode(encoding)
        except UnicodeDecodeError:
            continue
    return None


def _read_unknown(path: str) -> str:
    """未知扩展名：按纯文本尝试；嗅探为二进制则抛 BinaryFileError。"""
    with open(path, "rb") as f:
        head = f.read(8192)
        is_utf16 = head.startswith(_UTF16_BOMS)
        if b"\x00" in head and not is_utf16:
            raise BinaryFileError(f"二进制文件 {path}，无法提取文本")
        if os.path.getsize(path) > _UNKNOWN_TEXT_MAX:
            raise BinaryFileError(f"文件过大，不按文本解析：{path}")
        raw = head + f.read()
    text = _decode(raw)
    if text is None:
        raise BinaryFileError(f"无法识别文本编码：{path}")
    return text


# .txt/.md 的读取上限：超大日志整读进内存再 embedding 会 MemoryError 或持锁几小时
_TEXT_FILE_MAX = 20 * 1024 * 1024


def _read_text_file(path: str) -> str:
    if os.path.getsize(path) > _TEXT_FILE_MAX:
        raise ParseError(f"文本文件过大（超过 {_TEXT_FILE_MAX // (1024 * 1024)}MB），未解析全文")
    with open(path, "rb") as f:
        raw = f.read()
    text = _decode(raw)
    return text if text is not None else raw.decode("utf-8", errors="replace")


def _read_pdf(path: str) -> str:
    from pypdf import PdfReader

    reader = PdfReader(path)
    parts = []
    for page in reader.pages:
        try:
            text = page.extract_text()
        except Exception:
            text = None  # 提取失败的页跳过
        if text:
            parts.append(text)
    text = "\n\n".join(parts).strip()
    if not text:
        raise ParseError("无法提取文本（可能是扫描件）")
    return text


def _read_docx(path: str) -> str:
    import docx

    document = docx.Document(path)
    parts = [p.text for p in document.paragraphs if p.text and p.text.strip()]
    for table in document.tables:
        for row in table.rows:
            cells = [c.text.strip() for c in row.cells if c.text and c.text.strip()]
            if cells:
                parts.append(" | ".join(cells))
    text = "\n\n".join(parts).strip()
    if not text:
        raise ParseError("文档中没有可提取的文本")
    return text


def _read_xlsx(path: str) -> str:
    from openpyxl import load_workbook

    wb = load_workbook(path, read_only=True, data_only=True)
    parts = []
    try:  # read_only 模式持有文件句柄，异常时也要关，否则 Windows 下源文件一直被占用
        for sheet in wb.worksheets:
            rows = []
            for row in sheet.iter_rows(values_only=True):
                cells = [str(v).strip() for v in row if v is not None and str(v).strip()]
                if cells:
                    rows.append(f"{sheet.title} | {' | '.join(cells)}")
            if rows:
                parts.append("\n".join(rows))
    finally:
        wb.close()
    text = "\n\n".join(parts).strip()
    if not text:
        raise ParseError("表格中没有可提取的文本")
    return text


def _read_xls(path: str) -> str:
    """老式 Excel 97-2003（.xls）：xlrd 2.x 只读 .xls，输出格式与 _read_xlsx 一致。"""
    import xlrd

    try:
        wb = xlrd.open_workbook(path, on_demand=True)
    except xlrd.XLRDError as e:  # 加密、损坏，或其实是改了扩展名的 xlsx
        raise ParseError(f"无法读取 xls 文件：{e}") from e
    parts = []
    try:
        for sheet in wb.sheets():
            rows = []
            for r in range(sheet.nrows):
                cells = []
                for c in sheet.row(r):
                    v = c.value
                    if c.ctype == xlrd.XL_CELL_NUMBER and v == int(v):
                        v = int(v)  # xls 数字一律存成 float，整数去掉 ".0"
                    elif c.ctype == xlrd.XL_CELL_DATE:
                        v = xlrd.xldate_as_datetime(v, wb.datemode)
                    elif c.ctype in (xlrd.XL_CELL_EMPTY, xlrd.XL_CELL_BLANK, xlrd.XL_CELL_ERROR):
                        continue
                    s = str(v).strip()
                    if s:
                        cells.append(s)
                if cells:
                    rows.append(f"{sheet.name} | {' | '.join(cells)}")
            if rows:
                parts.append("\n".join(rows))
    finally:
        wb.release_resources()
    text = "\n\n".join(parts).strip()
    if not text:
        raise ParseError("表格中没有可提取的文本")
    return text


def _read_pptx(path: str) -> str:
    from pptx import Presentation

    prs = Presentation(path)
    parts = []
    for i, slide in enumerate(prs.slides, 1):
        texts = []
        for shape in slide.shapes:
            if shape.has_text_frame:
                t = "\n".join(
                    p.text for p in shape.text_frame.paragraphs if p.text.strip()
                )
                if t.strip():
                    texts.append(t.strip())
        if texts:
            parts.append(f"—— 第 {i} 页 ——\n" + "\n".join(texts))
    text = "\n\n".join(parts).strip()
    if not text:
        raise ParseError("演示文稿中没有可提取的文本")
    return text


_ocr_engine = None


def _get_ocr_engine():
    global _ocr_engine
    if _ocr_engine is None:
        try:
            from rapidocr_onnxruntime import RapidOCR
        except Exception as e:
            raise ParseError(f"OCR 组件不可用（{type(e).__name__}），请安装 rapidocr-onnxruntime")
        _ocr_engine = RapidOCR()
    return _ocr_engine


def _read_image_ocr(path: str) -> str:
    engine = _get_ocr_engine()
    try:
        result, _ = engine(path)
    except Exception as e:
        raise ParseError(f"OCR 识别失败：{e}")
    lines = [item[1].strip() for item in (result or []) if item[1] and item[1].strip()]
    text = "\n".join(lines).strip()
    if not text:
        raise ParseError("图片中没有识别到文字")
    return text
