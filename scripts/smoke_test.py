# -*- coding: utf-8 -*-
"""冒烟测试（不依赖 GUI）：临时样例文档走完整入库管道，再验证两个查询命中正确文档。

用法：python scripts/smoke_test.py
使用独立临时数据目录（PDA_DATA_DIR），不污染 ./data，结束后自动清理。
"""
import os
import shutil
import sys
import tempfile
from pathlib import Path

sys.stdout.reconfigure(encoding="utf-8", errors="replace")

# 在 import pda 之前设定独立数据目录与空配置文件
TMP_DIR = Path(tempfile.mkdtemp(prefix="pda_smoke_"))
os.environ["PDA_DATA_DIR"] = str(TMP_DIR / "data")
os.environ["PDA_CONFIG_PATH"] = str(TMP_DIR / "nonexistent_config.json")
os.environ.pop("PDA_API_KEY", None)  # 强制走"无 key"兜底路径

sys.path.insert(0, str(Path(__file__).resolve().parent.parent))

from pda import db, ingest, llm, qa  # noqa: E402


def make_sample_files(work_dir: Path) -> list:
    txt = work_dir / "张三的季度报告.txt"
    txt.write_text(
        "张三的季度报告\n\n"
        "Q3 季度销售业绩总结：本季度销售额达到 120 万元，同比增长 15%。"
        "主要增长来自华东区域的新客户拓展。\n\n"
        "下一季度目标：Q4 销售额冲刺 150 万元，重点推进大客户续约。",
        encoding="utf-8",
    )

    import docx

    document = docx.Document()
    document.add_heading("会议纪要", level=1)
    document.add_paragraph("时间：10 月 20 日，参会人：产品组全体成员。")
    document.add_paragraph("决议：产品发布会定在 11 月 15 日，地点为公司大会议室。")
    document.add_paragraph("会后行动项：市场部负责邀请函，研发部冻结功能。")
    docx_path = work_dir / "会议纪要.docx"
    document.save(str(docx_path))
    return [str(txt), str(docx_path)]


def check_hit(label: str, result: dict, expect_title: str, expect_keyword: str) -> bool:
    titles = [s["title"] for s in result["sources"]]
    joined = "\n".join(s["text"] for s in result["sources"])
    ok = any(expect_title in t for t in titles) and expect_keyword in joined
    print(f"  来源：{titles}")
    print(f"  命中检查：期望文档含『{expect_title}』且片段含『{expect_keyword}』 -> {'通过' if ok else '失败'}")
    return ok


def main():
    print(f"临时数据目录：{TMP_DIR}")
    work_dir = TMP_DIR / "samples"
    work_dir.mkdir(parents=True)
    files = make_sample_files(work_dir)

    print("\n=== 入库 ===")
    db.init_db()
    results = ingest.ingest_files(files)
    for r in results:
        if r["ok"]:
            print(f"  [OK] {r['title']}：doc_id={r['doc_id']}，{r['chunk_count']} 个片段")
        else:
            print(f"  [FAIL] {r['title']}：{r['error']}")
    assert all(r["ok"] for r in results), "入库失败"

    print(f"\nLLM 已配置：{llm.has_llm()}（预期 False，走兜底答案）")

    all_ok = True
    queries = [
        ("销售额是多少", "季度报告", "120 万"),
        ("发布会什么时候", "会议纪要", "11 月 15 日"),
    ]
    for question, expect_title, expect_keyword in queries:
        print(f"\n=== 查询：{question} ===")
        result = qa.answer(question)
        print(f"  答案：\n{result['answer']}")
        all_ok &= check_hit(question, result, expect_title, expect_keyword)

    shutil.rmtree(TMP_DIR, ignore_errors=True)
    print(f"\n临时目录已清理。总体结果：{'全部通过' if all_ok else '存在失败'}")
    sys.exit(0 if all_ok else 1)


if __name__ == "__main__":
    try:
        main()
    finally:
        # 中途异常也要清理临时目录（chroma 仍占用文件时 ignore_errors 兜底）
        shutil.rmtree(TMP_DIR, ignore_errors=True)
